"""Apply update costs to previous RA Rate card layout with validity splitting."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from load_input import (
    ROLE_UPDATE,
    _find_rate_card_header_row,
    load_update_dataframe,
)
from lane_columns import (
    DESTINATION_COUNTRY_COLUMN,
    ORIGIN_COUNTRY_COLUMN,
    add_lane_key,
    find_column,
    lane_key_from_values,
    normalize_port_name,
    resolve_header_lane_columns,
    resolve_rate_card_identity_columns,
)
from load_input import destination_country_code
from paths import OUTPUT_DIR, PREVIOUS_RA_DIR, PROCESSING_DIR, UPDATE_DIR

RATE_CARD_SHEET = "Rate card"
GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
VALUE_TOLERANCE = 0.009

VALID_FROM_COLUMN = "Valid from"
VALID_TO_COLUMN = "Valid to"
LANE_NUMBER_COLUMN = "Lane #"

UPDATE_VALID_FROM_ALIASES = ("RATE VALID FROM", "VALID FROM")
UPDATE_VALID_TO_ALIASES = ("RATE VALID TO", "VALID TO")

UPDATE_LSP_REFERENCE_ALIASES = ("CONTROLPAY LSP REFERENCE", "LSP REFERENCE")
UPDATE_SUPPLIER_NAME_ALIASES = ("SUPPLIER NAME",)
UPDATE_MODE_ALIASES = ("MODE",)
UPDATE_ORIGIN_REGION_ALIASES = ("ORIGIN REGION",)
UPDATE_DESTINATION_REGION_ALIASES = ("DESTINATION REGION",)
UPDATE_PUMA_DEST_REGION_ALIASES = ("PUMA DESTINATION REGION",)
UPDATE_ORIGIN_UNLOCODE_ALIASES = ("ORIGIN AIRPORT UNLOCODE", "ORIGIN PORT UNLOCODE")
UPDATE_DESTINATION_UNLOCODE_ALIASES = ("DESTINATION AIRPORT UNLOCODE", "DESTINATION PORT UNLOCODE")
UPDATE_ORIGIN_PORT_NAME_ALIASES = ("ORIGIN AIRPORT NAME", "ORIGIN PORT NAME")
UPDATE_DESTINATION_PORT_NAME_ALIASES = ("DESTINATION AIRPORT NAME", "DESTINATION PORT NAME")
UPDATE_ALL_IN_ALIASES = ("ALL IN (YES/NO)",)
UPDATE_CARRIER_PORTFOLIO_ALIASES = ("CARRIER PORTFOLIO",)
UPDATE_RATE_CARD_ALIASES = ("_source_file",)


@dataclass(frozen=True)
class CostTarget:
    update_column: str
    currency_column: str
    ra_pattern: str
    combine_baf_column: str | None = None
    subfield: str | None = None  # "min" | "per_kg" for AIR grouped costs


COST_TARGETS: tuple[CostTarget, ...] = (
    CostTarget("20' OTHC", "ORIGIN CURRENCY", r"origin terminal handling cost \(20ft\)"),
    CostTarget(
        "40' OTHC",
        "ORIGIN CURRENCY",
        r"origin terminal handling cost \((40ft|40hc|45hc)",
    ),
    CostTarget("CBL FEE (PER DOCUMENT)", "ORIGIN CURRENCY", r"bill of lading fee"),
    CostTarget("20' OCEAN RATE", "OCEAN CURRENCY", r"transport cost \(20ft\)", "20' BAF"),
    CostTarget("40' OCEAN RATE", "OCEAN CURRENCY", r"transport cost \(40ft\)", "40' BAF"),
    CostTarget("40' HC OCEAN RATE", "OCEAN CURRENCY", r"transport cost \(40hc\)", "40' HC BAF"),
    CostTarget("45' HC OCEAN RATE", "OCEAN CURRENCY", r"transport cost \(45hc\)", "45' HC BAF"),
    CostTarget("ETS RATE 20'", "OCEAN CURRENCY", r"eu ets \(emissions\) surcharge \(20ft"),
    CostTarget("ETS RATE 40'", "OCEAN CURRENCY", r"eu ets \(emissions\) surcharge \(40ft"),
    CostTarget("ETS RATE 40' HC", "OCEAN CURRENCY", r"eu ets \(emissions\) surcharge \(40hc"),
    CostTarget("ETS RATE 45' HC", "OCEAN CURRENCY", r"eu ets \(emissions\) surcharge \(45hc"),
    CostTarget("20' DTHC", "DESTINATION CURRENCY", r"destination terminal handling cost \(20ft\)"),
    CostTarget(
        "40' DTHC",
        "DESTINATION CURRENCY",
        r"destination terminal handling cost \((40ft|40hc|45hc)",
    ),
    CostTarget("LCL RATE PER CBM", "OCEAN CURRENCY", r"lcl rate per cbm|origin terminal handling cost per cbm"),
    CostTarget("ETS RATE LCL", "OCEAN CURRENCY", r"eu ets.*lcl|ets.*per cbm"),
    # AIR (per-kg / MIN grouped costs)
    CostTarget(
        "MIN OTHC (PER HAWB)",
        "ORIGIN CURRENCY",
        r"^origin terminal handling cost$",
        subfield="min",
    ),
    CostTarget(
        "OTHC PER KG VOL",
        "ORIGIN CURRENCY",
        r"^origin terminal handling cost$",
        subfield="per_kg",
    ),
    CostTarget("MIN (HAWB)", "AIRFREIGHT CURRENCY", r"^transport cost$", subfield="min"),
    CostTarget(
        "FLAT RATE PER KG VOL",
        "AIRFREIGHT CURRENCY",
        r"^transport cost$",
        subfield="per_kg",
    ),
    CostTarget(
        "MIN DTHC (PER HAWB)",
        "DESTINATION CURRENCY",
        r"^destination terminal handling cost$",
        subfield="min",
    ),
    CostTarget(
        "DTHC PER KG VOL",
        "DESTINATION CURRENCY",
        r"^destination terminal handling cost$",
        subfield="per_kg",
    ),
    CostTarget(
        "MIN ON CARRIAGE (PER HAWB) (APPLICABLE FOR ALL DELIVERIES)",
        "DESTINATION CURRENCY",
        r"^on carriage$",
        subfield="min",
    ),
    CostTarget(
        "ON CARRIAGE PER KG VOL (APPLICABLE FOR ALL DELIVERIES)",
        "DESTINATION CURRENCY",
        r"^on carriage$",
        subfield="per_kg",
    ),
    CostTarget(
        "ON CARRIAGE FTL (APPLICABLE FOR ALL DELIVERIES)",
        "DESTINATION CURRENCY",
        r"^on carriage$",
        subfield="per_shipment",
    ),
)


@dataclass
class CostColumn:
    label: str
    currency_col: int
    value_col: int
    min_value_col: int | None = None
    per_shipment_col: int | None = None
    supplier: str | None = None


@dataclass(frozen=True)
class SeaCostApplication:
    ra_pattern: str
    currency: object
    amount: object


KNOWN_SUPPLIER_TOKENS: tuple[str, ...] = (
    "DSV",
    "MAERSK",
    "MSC",
    "CMA",
    "HAPAG",
    "ONE",
    "EVERGREEN",
    "COSCO",
    "JAS",
    "KUEHNE",
    "KN",
    "HMM",
    "YANG MING",
    "ZIM",
    "OOCL",
    "PIL",
)

SEA_TRANSPORT_RA_PATTERNS: dict[str, str] = {
    "20ft": r"transport cost \(20ft\)",
    "40ft": r"transport cost \(40ft\)",
    "40hc": r"transport cost \(40hc\)",
    "45hc": r"transport cost \(45hc\)",
}

SEA_OTHC_RA_PATTERNS: dict[str, str] = {
    "20ft": r"origin terminal handling cost \(20ft\)",
    "40ft": r"origin terminal handling cost \(40ft\)",
    "40hc": r"origin terminal handling cost \(40hc\)",
    "45hc": r"origin terminal handling cost \(45hc\)",
}

SEA_DTHC_RA_PATTERNS: dict[str, str] = {
    "20ft": r"destination terminal handling cost \(20ft\)",
    "40ft": r"destination terminal handling cost \(40ft\)",
    "40hc": r"destination terminal handling cost \(40hc\)",
    "45hc": r"destination terminal handling cost \(45hc\)",
}

SEA_ETS_RA_PATTERNS: dict[str, str] = {
    "20ft": r"eu ets \(emissions\) surcharge \(20ft",
    "40ft": r"eu ets \(emissions\) surcharge \(40ft",
    "40hc": r"eu ets \(emissions\) surcharge \(40hc",
    "45hc": r"eu ets \(emissions\) surcharge \(45hc",
}

AIR_COST_TARGETS: tuple[CostTarget, ...] = tuple(
    target for target in COST_TARGETS if target.subfield is not None
)


def normalize_label(value: object) -> str:
    text = str(value or "").replace("\r\n", " ").strip().lower()
    return re.sub(r"\s+", " ", text)


def supplier_token_from_text(value: object) -> str | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    text = str(value).strip().upper()
    if not text:
        return None
    for token in KNOWN_SUPPLIER_TOKENS:
        if token in text:
            return token.replace(" ", "")
    compact = re.sub(r"[^A-Z0-9]", "", text)
    if 2 <= len(compact) <= 8 and compact.isalpha():
        return compact
    return None


def read_column_supplier(
    ws: Worksheet,
    label_row: int,
    start_col: int,
    end_col: int,
) -> str | None:
    label_supplier = supplier_token_from_text(ws.cell(label_row, start_col).value)
    if label_supplier:
        return label_supplier
    for row_idx in range(max(1, label_row - 3), label_row):
        for col in range(start_col, end_col + 1):
            supplier = supplier_token_from_text(ws.cell(row_idx, col).value)
            if supplier:
                return supplier
    return None


def suppliers_match(column_supplier: str | None, update_supplier: str | None) -> bool:
    if column_supplier is None:
        return True
    if update_supplier is None:
        return False
    return column_supplier == update_supplier


def update_column_has_value(update_row: pd.Series, column_name: str) -> bool:
    if column_name not in update_row.index:
        return False
    return pd.notna(update_row[column_name])


def update_numeric_value(update_row: pd.Series, column_name: str) -> float | None:
    if not update_column_has_value(update_row, column_name):
        return None
    try:
        return float(update_row[column_name])
    except (TypeError, ValueError):
        return None


def update_currency_value(update_row: pd.Series, column_name: str) -> object | None:
    if column_name not in update_row.index:
        return None
    currency = update_row[column_name]
    if pd.isna(currency):
        return None
    return currency


def combine_transport_amount(update_row: pd.Series, ocean_column: str, baf_column: str) -> float | None:
    ocean_amount = update_numeric_value(update_row, ocean_column)
    if ocean_amount is None:
        return None
    baf_amount = update_numeric_value(update_row, baf_column)
    if baf_amount is not None:
        return ocean_amount + baf_amount
    return ocean_amount


def split_40_size_amounts(
    update_row: pd.Series,
    *,
    generic_column: str,
    hc_column: str | None = None,
    hc45_column: str | None = None,
) -> dict[str, float]:
    amounts: dict[str, float] = {}
    has_generic = update_column_has_value(update_row, generic_column)
    has_40hc = hc_column is not None and update_column_has_value(update_row, hc_column)
    has_45hc = hc45_column is not None and update_column_has_value(update_row, hc45_column)

    if has_generic:
        generic_amount = float(update_row[generic_column])
        if has_40hc or has_45hc:
            amounts["40ft"] = generic_amount
        else:
            amounts["40ft"] = generic_amount
            amounts["40hc"] = generic_amount
            amounts["45hc"] = generic_amount

    if has_40hc and hc_column is not None:
        amounts["40hc"] = float(update_row[hc_column])
    if has_45hc and hc45_column is not None:
        amounts["45hc"] = float(update_row[hc45_column])
    return amounts


def split_40_transport_amounts(update_row: pd.Series) -> dict[str, float]:
    amounts: dict[str, float] = {}
    has_generic = update_column_has_value(update_row, "40' OCEAN RATE")
    has_40hc = update_column_has_value(update_row, "40' HC OCEAN RATE")
    has_45hc = update_column_has_value(update_row, "45' HC OCEAN RATE")

    if has_generic:
        ocean_amount = update_numeric_value(update_row, "40' OCEAN RATE")
        if ocean_amount is None:
            ocean_amount = 0.0
        baf_40 = update_numeric_value(update_row, "40' BAF") or 0.0
        if has_40hc or has_45hc:
            amounts["40ft"] = ocean_amount + baf_40
        else:
            baf_40hc = update_numeric_value(update_row, "40' HC BAF")
            baf_45hc = update_numeric_value(update_row, "45' HC BAF")
            amounts["40ft"] = ocean_amount + baf_40
            amounts["40hc"] = ocean_amount + (baf_40hc if baf_40hc is not None else baf_40)
            amounts["45hc"] = ocean_amount + (baf_45hc if baf_45hc is not None else baf_40)

    if has_40hc:
        combined = combine_transport_amount(update_row, "40' HC OCEAN RATE", "40' HC BAF")
        if combined is not None:
            amounts["40hc"] = combined
    if has_45hc:
        combined = combine_transport_amount(update_row, "45' HC OCEAN RATE", "45' HC BAF")
        if combined is not None:
            amounts["45hc"] = combined
    return amounts


def add_sea_size_applications(
    applications: list[SeaCostApplication],
    *,
    update_row: pd.Series,
    currency_column: str,
    ra_patterns: dict[str, str],
    amounts_by_size: dict[str, float],
) -> None:
    currency = update_currency_value(update_row, currency_column)
    if currency is None:
        return
    for size_key, amount in amounts_by_size.items():
        pattern = ra_patterns.get(size_key)
        if pattern is None:
            continue
        applications.append(SeaCostApplication(pattern, currency, amount))


def sea_cost_applications(update_row: pd.Series) -> list[SeaCostApplication]:
    applications: list[SeaCostApplication] = []

    if update_column_has_value(update_row, "20' OTHC"):
        add_sea_size_applications(
            applications,
            update_row=update_row,
            currency_column="ORIGIN CURRENCY",
            ra_patterns=SEA_OTHC_RA_PATTERNS,
            amounts_by_size={"20ft": float(update_row["20' OTHC"])},
        )
    add_sea_size_applications(
        applications,
        update_row=update_row,
        currency_column="ORIGIN CURRENCY",
        ra_patterns=SEA_OTHC_RA_PATTERNS,
        amounts_by_size=split_40_size_amounts(update_row, generic_column="40' OTHC"),
    )

    if update_column_has_value(update_row, "20' OCEAN RATE"):
        combined = combine_transport_amount(update_row, "20' OCEAN RATE", "20' BAF")
        if combined is not None:
            currency = update_currency_value(update_row, "OCEAN CURRENCY")
            if currency is not None:
                applications.append(
                    SeaCostApplication(SEA_TRANSPORT_RA_PATTERNS["20ft"], currency, combined)
                )
    add_sea_size_applications(
        applications,
        update_row=update_row,
        currency_column="OCEAN CURRENCY",
        ra_patterns=SEA_TRANSPORT_RA_PATTERNS,
        amounts_by_size=split_40_transport_amounts(update_row),
    )

    if update_column_has_value(update_row, "ETS RATE 20'"):
        currency = update_currency_value(update_row, "OCEAN CURRENCY")
        if currency is not None:
            applications.append(
                SeaCostApplication(
                    SEA_ETS_RA_PATTERNS["20ft"],
                    currency,
                    float(update_row["ETS RATE 20'"]),
                )
            )
    add_sea_size_applications(
        applications,
        update_row=update_row,
        currency_column="OCEAN CURRENCY",
        ra_patterns=SEA_ETS_RA_PATTERNS,
        amounts_by_size=split_40_size_amounts(
            update_row,
            generic_column="ETS RATE 40'",
            hc_column="ETS RATE 40' HC",
            hc45_column="ETS RATE 45' HC",
        ),
    )

    if update_column_has_value(update_row, "20' DTHC"):
        add_sea_size_applications(
            applications,
            update_row=update_row,
            currency_column="DESTINATION CURRENCY",
            ra_patterns=SEA_DTHC_RA_PATTERNS,
            amounts_by_size={"20ft": float(update_row["20' DTHC"])},
        )
    add_sea_size_applications(
        applications,
        update_row=update_row,
        currency_column="DESTINATION CURRENCY",
        ra_patterns=SEA_DTHC_RA_PATTERNS,
        amounts_by_size=split_40_size_amounts(update_row, generic_column="40' DTHC"),
    )

    if update_column_has_value(update_row, "CBL FEE (PER DOCUMENT)"):
        currency = update_currency_value(update_row, "ORIGIN CURRENCY")
        if currency is not None:
            applications.append(
                SeaCostApplication(
                    r"bill of lading fee",
                    currency,
                    update_row["CBL FEE (PER DOCUMENT)"],
                )
            )

    if update_column_has_value(update_row, "LCL RATE PER CBM"):
        currency = update_currency_value(update_row, "OCEAN CURRENCY")
        if currency is not None:
            applications.append(
                SeaCostApplication(
                    r"lcl rate per cbm|origin terminal handling cost per cbm",
                    currency,
                    update_row["LCL RATE PER CBM"],
                )
            )

    if update_column_has_value(update_row, "ETS RATE LCL"):
        currency = update_currency_value(update_row, "OCEAN CURRENCY")
        if currency is not None:
            applications.append(
                SeaCostApplication(
                    r"eu ets.*lcl|ets.*per cbm",
                    currency,
                    update_row["ETS RATE LCL"],
                )
            )

    return applications


def is_air_cost_layout(cost_columns: list[CostColumn]) -> bool:
    return any(column.min_value_col is not None for column in cost_columns)


def parse_date(value: object) -> date | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, pd.Timestamp):
        return value.date()
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    parsed = pd.to_datetime(value, dayfirst=True, errors="coerce")
    if pd.isna(parsed):
        return None
    return parsed.date()


def format_ra_date(value: date) -> str:
    return value.strftime("%d.%m.%Y")


def numeric_equal(left: object, right: object) -> bool:
    if left is None and right is None:
        return True
    if left is None or right is None:
        return False
    try:
        return abs(float(left) - float(right)) <= VALUE_TOLERANCE
    except (TypeError, ValueError):
        return str(left).strip() == str(right).strip()


def split_validity_periods(
    update_start: date,
    update_end: date,
    ra_periods: list[tuple[date, date]],
) -> list[tuple[date, date]]:
    boundaries = {update_start, update_end + timedelta(days=1)}
    for ra_start, ra_end in ra_periods:
        boundaries.add(ra_start)
        boundaries.add(ra_end + timedelta(days=1))

    ordered = sorted(
        boundary
        for boundary in boundaries
        if update_start <= boundary <= update_end + timedelta(days=1)
    )
    segments: list[tuple[date, date]] = []
    for index in range(len(ordered) - 1):
        segment_start = ordered[index]
        segment_end = ordered[index + 1] - timedelta(days=1)
        if segment_start > segment_end:
            continue
        clipped_start = max(segment_start, update_start)
        clipped_end = min(segment_end, update_end)
        if clipped_start <= clipped_end:
            segments.append((clipped_start, clipped_end))
    return segments


def lane_id_value(value: object) -> int | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    parsed = pd.to_numeric(str(value).replace(" ", ""), errors="coerce")
    if pd.isna(parsed):
        return None
    return int(parsed)


def detect_cost_label_row(ws: Worksheet) -> int:
    for row_idx in range(2, 8):
        for col in range(1, ws.max_column + 1):
            label = normalize_label(ws.cell(row_idx, col).value)
            if not label or "applies if" in label:
                continue
            if label in {"origin terminal handling cost", "transport cost"}:
                return row_idx
            if label.startswith("origin terminal handling cost ("):
                return row_idx
    return 4


def find_currency_subheader_row(ws: Worksheet, start_col: int, label_row: int) -> int:
    for row_idx in range(label_row + 1, label_row + 6):
        if normalize_label(ws.cell(row_idx, start_col).value) == "currency":
            return row_idx
    return label_row + 4


def parse_merged_cost_group(
    ws: Worksheet,
    label_row: int,
    start_col: int,
    end_col: int,
) -> CostColumn | None:
    label = normalize_label(ws.cell(label_row, start_col).value)
    if not label or "applies if" in label:
        return None

    subheader_row = find_currency_subheader_row(ws, start_col, label_row)
    min_row = subheader_row - 1
    min_value_col: int | None = None
    per_shipment_col: int | None = None
    per_kg_col: int | None = None

    for col in range(start_col + 1, end_col + 1):
        subheader = normalize_label(ws.cell(subheader_row, col).value)
        min_label = normalize_label(ws.cell(min_row, col).value)
        if min_label == "min" or subheader == "flat":
            min_value_col = col
        elif subheader in {"p/unit", "per unit"}:
            per_kg_col = col
        elif subheader == "per shipment":
            per_shipment_col = col

    value_col = per_kg_col or per_shipment_col or (start_col + 1 if end_col > start_col else start_col)
    return CostColumn(
        label=label,
        currency_col=start_col,
        value_col=value_col,
        min_value_col=min_value_col,
        per_shipment_col=per_shipment_col if per_shipment_col != value_col else None,
        supplier=read_column_supplier(ws, label_row, start_col, end_col),
    )


def parse_cost_columns(ws: Worksheet) -> list[CostColumn]:
    label_row = detect_cost_label_row(ws)
    merged_ranges = sorted(
        (
            merged
            for merged in ws.merged_cells.ranges
            if merged.min_row == label_row and merged.max_row == label_row
        ),
        key=lambda merged: merged.min_col,
    )
    if merged_ranges:
        columns: list[CostColumn] = []
        for merged in merged_ranges:
            column = parse_merged_cost_group(ws, label_row, merged.min_col, merged.max_col)
            if column is not None:
                columns.append(column)
        return columns

    columns = []
    for col in range(1, ws.max_column):
        label = ws.cell(label_row, col).value
        if label is None or not str(label).strip():
            continue
        columns.append(
            CostColumn(
                label=normalize_label(label),
                currency_col=col,
                value_col=col + 1,
                supplier=read_column_supplier(ws, label_row, col, col),
            )
        )
    return columns


def target_value_column(column: CostColumn, subfield: str | None) -> int | None:
    if subfield == "min":
        return column.min_value_col
    if subfield == "per_kg":
        return column.value_col if column.min_value_col is not None else column.value_col
    if subfield == "per_shipment":
        return column.per_shipment_col or column.value_col
    return column.value_col


def header_column_map(ws: Worksheet, header_row: int) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(header_row, col).value
        if value is not None and str(value).strip():
            key = str(value).strip()
            if key not in mapping:
                mapping[key] = col
    return mapping


def row_values(ws: Worksheet, row_idx: int, max_col: int) -> list[object]:
    return [ws.cell(row_idx, col).value for col in range(1, max_col + 1)]


def write_row_values(ws: Worksheet, row_idx: int, values: list[object]) -> None:
    for col, value in enumerate(values, start=1):
        ws.cell(row_idx, col).value = value


def highlight_cell(ws: Worksheet, row_idx: int, col_idx: int, fill: PatternFill = GREEN_FILL) -> None:
    ws.cell(row_idx, col_idx).fill = fill


def highlight_shipment_details_block(
    ws: Worksheet, row_idx: int, last_col: int, fill: PatternFill = GREEN_FILL
) -> None:
    for col in range(1, last_col + 1):
        highlight_cell(ws, row_idx, col, fill)


def highlight_new_lane_row(
    ws: Worksheet,
    row_idx: int,
    *,
    shipment_last_col: int,
    cost_columns: list[CostColumn],
    valid_from_col: int | None,
    valid_to_col: int | None,
) -> None:
    highlight_shipment_details_block(ws, row_idx, shipment_last_col, YELLOW_FILL)
    if valid_from_col is not None:
        highlight_cell(ws, row_idx, valid_from_col, YELLOW_FILL)
    if valid_to_col is not None:
        highlight_cell(ws, row_idx, valid_to_col, YELLOW_FILL)
    for column in cost_columns:
        highlight_cell(ws, row_idx, column.currency_col, YELLOW_FILL)
        highlight_cell(ws, row_idx, column.value_col, YELLOW_FILL)
        if column.min_value_col is not None:
            highlight_cell(ws, row_idx, column.min_value_col, YELLOW_FILL)
        if column.per_shipment_col is not None:
            highlight_cell(ws, row_idx, column.per_shipment_col, YELLOW_FILL)


def shipment_details_last_column(cost_columns: list[CostColumn]) -> int:
    if not cost_columns:
        return 25
    return min(column.currency_col for column in cost_columns) - 1


def update_validity_dates(update_row: pd.Series) -> tuple[date | None, date | None]:
    valid_from_column = find_column(update_row.index, UPDATE_VALID_FROM_ALIASES)
    valid_to_column = find_column(update_row.index, UPDATE_VALID_TO_ALIASES)
    if valid_from_column is None or valid_to_column is None:
        return None, None
    return parse_date(update_row[valid_from_column]), parse_date(update_row[valid_to_column])


def update_cost_amount(update_row: pd.Series, target: CostTarget) -> tuple[object, object] | None:
    if target.update_column not in update_row.index:
        return None
    amount = update_row[target.update_column]
    if pd.isna(amount):
        return None
    if target.combine_baf_column and target.combine_baf_column in update_row.index:
        baf = update_row[target.combine_baf_column]
        if pd.notna(baf):
            amount = float(amount) + float(baf)
    currency = update_row.get(target.currency_column)
    if pd.isna(currency):
        currency = None
    return currency, amount


def apply_cost_updates(
    ws: Worksheet,
    row_idx: int,
    update_row: pd.Series,
    cost_columns: list[CostColumn],
    *,
    highlight: bool = True,
) -> int:
    if is_air_cost_layout(cost_columns):
        return apply_air_cost_updates(ws, row_idx, update_row, cost_columns, highlight=highlight)
    return apply_sea_cost_updates(ws, row_idx, update_row, cost_columns, highlight=highlight)


def apply_air_cost_updates(
    ws: Worksheet,
    row_idx: int,
    update_row: pd.Series,
    cost_columns: list[CostColumn],
    *,
    highlight: bool = True,
) -> int:
    changed_cells = 0
    for target in AIR_COST_TARGETS:
        update_value = update_cost_amount(update_row, target)
        if update_value is None:
            continue
        currency, amount = update_value
        pattern = re.compile(target.ra_pattern)
        for column in cost_columns:
            if not pattern.search(column.label):
                continue
            value_col = target_value_column(column, target.subfield)
            if value_col is None:
                continue
            old_currency = ws.cell(row_idx, column.currency_col).value
            old_amount = ws.cell(row_idx, value_col).value
            ws.cell(row_idx, column.currency_col).value = currency
            ws.cell(row_idx, value_col).value = amount
            if highlight and (
                not numeric_equal(old_currency, currency) or not numeric_equal(old_amount, amount)
            ):
                highlight_cell(ws, row_idx, column.currency_col)
                highlight_cell(ws, row_idx, value_col)
                changed_cells += 2
            elif not highlight:
                changed_cells += 2
    return changed_cells


def apply_sea_cost_updates(
    ws: Worksheet,
    row_idx: int,
    update_row: pd.Series,
    cost_columns: list[CostColumn],
    *,
    highlight: bool = True,
) -> int:
    changed_cells = 0
    update_supplier = update_supplier_token(update_row)
    for application in sea_cost_applications(update_row):
        pattern = re.compile(application.ra_pattern)
        for column in cost_columns:
            if not pattern.search(column.label):
                continue
            if not suppliers_match(column.supplier, update_supplier):
                continue
            old_currency = ws.cell(row_idx, column.currency_col).value
            old_amount = ws.cell(row_idx, column.value_col).value
            ws.cell(row_idx, column.currency_col).value = application.currency
            ws.cell(row_idx, column.value_col).value = application.amount
            if highlight and (
                not numeric_equal(old_currency, application.currency)
                or not numeric_equal(old_amount, application.amount)
            ):
                highlight_cell(ws, row_idx, column.currency_col)
                highlight_cell(ws, row_idx, column.value_col)
                changed_cells += 2
            elif not highlight:
                changed_cells += 2
    return changed_cells


def choose_template_row(
    segment: tuple[date, date],
    lane_rows: list[tuple[tuple[date, date], list[object]]],
) -> list[object]:
    segment_start, segment_end = segment
    for period, values in lane_rows:
        if period[0] <= segment_end and period[1] >= segment_start:
            return list(values)
    return list(lane_rows[0][1])


def load_update_rows(update_path: Path | None = None) -> pd.DataFrame:
    if update_path is not None and update_path.exists():
        workbook = pd.ExcelFile(update_path)
        frames = [pd.read_excel(workbook, sheet_name=name) for name in workbook.sheet_names]
        update = pd.concat(frames, ignore_index=True, sort=False)
    elif (PROCESSING_DIR / f"{ROLE_UPDATE}.xlsx").exists():
        workbook = pd.ExcelFile(PROCESSING_DIR / f"{ROLE_UPDATE}.xlsx")
        frames = [pd.read_excel(workbook, sheet_name=name) for name in workbook.sheet_names]
        update = pd.concat(frames, ignore_index=True, sort=False)
    else:
        update = load_update_dataframe(list(UPDATE_DIR.glob("*.xlsx"))[0])
    return add_lane_key(update)


def build_update_lookup(update: pd.DataFrame) -> dict[str, pd.Series]:
    lookup: dict[str, pd.Series] = {}
    for _, row in update.iterrows():
        key = row.get("_lane_key")
        if key and key not in lookup:
            lookup[str(key)] = row
    return lookup


def update_field_value(update_row: pd.Series, aliases: tuple[str, ...]) -> object | None:
    column = find_column(update_row.index, aliases)
    if column is None:
        return None
    value = update_row[column]
    if pd.isna(value):
        return None
    return value


def update_supplier_token(update_row: pd.Series) -> str | None:
    for aliases in (
        UPDATE_SUPPLIER_NAME_ALIASES,
        ("SUPPLIER CODE",),
        UPDATE_LSP_REFERENCE_ALIASES,
        ("CORE CARRIER",),
    ):
        value = update_field_value(update_row, aliases)
        supplier = supplier_token_from_text(value)
        if supplier:
            return supplier
    return None


def map_transport_mode(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip().upper()
    if "AIR" in text:
        return "AIR"
    if "SEA" in text:
        return "SEA"
    return str(value).strip()


def rate_card_name_from_update(update_row: pd.Series) -> str | None:
    source = update_field_value(update_row, UPDATE_RATE_CARD_ALIASES)
    if source is None:
        return None
    return Path(str(source)).stem


def ra_lane_keys_from_data(
    data_rows: list[list[object]],
    columns: dict[str, int],
    lane_columns: dict[str, str],
) -> set[str]:
    keys: set[str] = set()
    for values in data_rows:
        lane_key = lane_key_from_values(values, columns, lane_columns)
        if lane_key:
            keys.add(lane_key)
    return keys


def ra_lane_ids_from_data(
    data_rows: list[list[object]],
    columns: dict[str, int],
    lane_id_column: str,
) -> set[int]:
    ids: set[int] = set()
    for values in data_rows:
        lane_id = lane_id_value(values[columns[lane_id_column] - 1])
        if lane_id is not None:
            ids.add(lane_id)
    return ids


def ra_destination_codes_from_data(
    data_rows: list[list[object]],
    columns: dict[str, int],
    destination_country_column: str,
) -> set[str]:
    codes: set[str] = set()
    for values in data_rows:
        code = destination_country_code(values[columns[destination_country_column] - 1])
        if code:
            codes.add(code)
    return codes


def unmapped_update_rows_for_ra(
    update_lookup: dict[str, pd.Series],
    ra_lane_keys: set[str],
    ra_lane_ids: set[int],
    ra_destination_codes: set[str],
) -> list[pd.Series]:
    rows: list[pd.Series] = []
    for lane_key, update_row in update_lookup.items():
        if lane_key in ra_lane_keys:
            continue
        lane_id = lane_id_value(update_row.get("LANE ID"))
        if lane_id is not None and lane_id in ra_lane_ids:
            continue
        destination_code = destination_country_code(update_row.get(DESTINATION_COUNTRY_COLUMN))
        if destination_code is None or destination_code not in ra_destination_codes:
            continue
        rows.append(update_row)
    rows.sort(
        key=lambda row: (
            destination_country_code(row.get(DESTINATION_COUNTRY_COLUMN)) or "",
            destination_country_code(row.get(ORIGIN_COUNTRY_COLUMN)) or "",
            str(update_field_value(row, UPDATE_ORIGIN_PORT_NAME_ALIASES) or ""),
            str(update_field_value(row, UPDATE_DESTINATION_PORT_NAME_ALIASES) or ""),
            lane_id_value(row.get("LANE ID")) or 0,
        )
    )
    return rows


def choose_new_lane_template(
    data_rows: list[list[object]],
    update_row: pd.Series,
    columns: dict[str, int],
    lane_columns: dict[str, str],
) -> list[object]:
    update_dest = destination_country_code(update_row.get(DESTINATION_COUNTRY_COLUMN))
    update_origin = destination_country_code(update_row.get(ORIGIN_COUNTRY_COLUMN))
    for values in reversed(data_rows):
        dest = destination_country_code(values[columns[lane_columns["destination_country"]] - 1])
        origin = destination_country_code(values[columns[lane_columns["origin_country"]] - 1])
        if dest == update_dest and (update_origin is None or origin == update_origin):
            return list(values)
    return list(data_rows[-1])


def set_row_column(values: list[object], columns: dict[str, int], header_name: str, value: object) -> None:
    column = columns.get(header_name)
    if column is None:
        return
    values[column - 1] = value


def build_ra_row_from_update(
    template_values: list[object],
    update_row: pd.Series,
    *,
    columns: dict[str, int],
    lane_columns: dict[str, str],
    identity_columns: dict[str, str],
    lane_number: int | None,
    rate_card_name: str | None,
) -> list[object]:
    row_copy = list(template_values)
    max_col = len(columns) and max(columns.values()) or len(row_copy)
    if len(row_copy) < max_col:
        row_copy.extend([None] * (max_col - len(row_copy)))

    if lane_number is not None and "lane_number" in identity_columns:
        set_row_column(row_copy, columns, identity_columns["lane_number"], lane_number)
    set_row_column(row_copy, columns, "Rate Card", rate_card_name)
    set_row_column(row_copy, columns, identity_columns["lane_id"], update_row.get("LANE ID"))
    set_row_column(row_copy, columns, "LSP REFERENCE", update_field_value(update_row, UPDATE_LSP_REFERENCE_ALIASES))
    set_row_column(
        row_copy,
        columns,
        "Carrier Name",
        update_field_value(update_row, UPDATE_LSP_REFERENCE_ALIASES)
        or update_field_value(update_row, UPDATE_SUPPLIER_NAME_ALIASES),
    )
    set_row_column(row_copy, columns, "Transport Mode", map_transport_mode(update_field_value(update_row, UPDATE_MODE_ALIASES)))
    set_row_column(row_copy, columns, "Origin region", update_field_value(update_row, UPDATE_ORIGIN_REGION_ALIASES))
    set_row_column(row_copy, columns, lane_columns["origin_country"], update_row.get(ORIGIN_COUNTRY_COLUMN))
    set_row_column(row_copy, columns, "Destination region", update_field_value(update_row, UPDATE_DESTINATION_REGION_ALIASES))
    set_row_column(row_copy, columns, "Puma Destination region", update_field_value(update_row, UPDATE_PUMA_DEST_REGION_ALIASES))
    set_row_column(row_copy, columns, lane_columns["destination_country"], update_row.get(DESTINATION_COUNTRY_COLUMN))
    set_row_column(row_copy, columns, "ORIGIN AIRPORT UNLOCODE", update_field_value(update_row, UPDATE_ORIGIN_UNLOCODE_ALIASES))
    set_row_column(
        row_copy,
        columns,
        "DESTINATION AIRPORT UNLOCODE",
        update_field_value(update_row, UPDATE_DESTINATION_UNLOCODE_ALIASES),
    )
    set_row_column(
        row_copy,
        columns,
        lane_columns["origin_port"],
        update_field_value(update_row, ("ORIGIN AIRPORT", "ORIGIN PORT")),
    )
    set_row_column(row_copy, columns, "ORIGIN AIRPORT NAME", update_field_value(update_row, UPDATE_ORIGIN_PORT_NAME_ALIASES))
    set_row_column(
        row_copy,
        columns,
        lane_columns["destination_port"],
        update_field_value(update_row, ("DESTINATION AIRPORT", "DESTINATION PORT")),
    )
    set_row_column(
        row_copy,
        columns,
        "DESTINATION AIRPORT NAME",
        update_field_value(update_row, UPDATE_DESTINATION_PORT_NAME_ALIASES),
    )
    set_row_column(row_copy, columns, "ALL IN (YES/NO)", update_field_value(update_row, UPDATE_ALL_IN_ALIASES))
    set_row_column(row_copy, columns, "CARRIER PORTFOLIO", update_field_value(update_row, UPDATE_CARRIER_PORTFOLIO_ALIASES))

    update_start, update_end = update_validity_dates(update_row)
    if update_start is not None:
        set_row_column(row_copy, columns, identity_columns["valid_from"], format_ra_date(update_start))
    if update_end is not None:
        set_row_column(row_copy, columns, identity_columns["valid_to"], format_ra_date(update_end))
    return row_copy


def append_unmapped_lanes(
    data_rows: list[list[object]],
    update_lookup: dict[str, pd.Series],
    *,
    columns: dict[str, int],
    lane_columns: dict[str, str],
    identity_columns: dict[str, str],
    next_lane_number: int | None,
) -> tuple[list[list[object]], int, dict[str, int]]:
    ra_lane_keys = ra_lane_keys_from_data(data_rows, columns, lane_columns)
    ra_lane_ids = ra_lane_ids_from_data(data_rows, columns, identity_columns["lane_id"])
    ra_destination_codes = ra_destination_codes_from_data(
        data_rows,
        columns,
        lane_columns["destination_country"],
    )
    unmapped_rows = unmapped_update_rows_for_ra(
        update_lookup,
        ra_lane_keys,
        ra_lane_ids,
        ra_destination_codes,
    )
    stats = {"lanes_added": 0, "rows_added_unmapped": 0}
    if not unmapped_rows:
        return data_rows, next_lane_number or 0, stats

    rate_card_name = rate_card_name_from_update(unmapped_rows[0])
    if rate_card_name is None:
        for values in data_rows:
            existing = columns.get("Rate Card")
            if existing is not None and values[existing - 1]:
                rate_card_name = values[existing - 1]
                break

    appended: list[list[object]] = []
    for update_row in unmapped_rows:
        template_values = choose_new_lane_template(data_rows, update_row, columns, lane_columns)
        lane_number = next_lane_number
        if lane_number is not None:
            next_lane_number = lane_number + 1
        appended.append(
            build_ra_row_from_update(
                template_values,
                update_row,
                columns=columns,
                lane_columns=lane_columns,
                identity_columns=identity_columns,
                lane_number=lane_number,
                rate_card_name=rate_card_name,
            )
        )
        stats["lanes_added"] += 1
        stats["rows_added_unmapped"] += 1

    return data_rows + appended, next_lane_number or 0, stats


def is_air_rate_card(ws: Worksheet) -> bool:
    header_row = _find_rate_card_header_row_from_sheet(ws)
    columns = header_column_map(ws, header_row)
    header_keys = {str(name).strip().upper() for name in columns}

    if "MODE" in header_keys or "SERVICE TYPE" in header_keys:
        return False
    if "LANE ID" in header_keys or "TRANSPORT MODE" in header_keys:
        return True

    label_row = detect_cost_label_row(ws)
    for col in range(1, min(ws.max_column, 120) + 1):
        label = normalize_label(ws.cell(label_row, col).value)
        if label.startswith("origin terminal handling cost ("):
            return False
        if label in {"origin terminal handling cost", "transport cost"}:
            return True
    return False


AIR_SPLIT_KEY_ORDER: tuple[str, ...] = (
    "HK|US|HKG|LAX",
    "CN|US|CAN|LAX",
    "CN|US|XMN|LAX",
    "CN|US|SZX|LAX",
    "CN|US|PVG|ORD",
    "CN|US|CGO|ORD",
    "HK|US|HKG|ORD",
    "CN|US|CTU|LAX",
)


def air_split_sort_key(lane_key: str, lane_number: object) -> tuple[int, int]:
    try:
        group_index = AIR_SPLIT_KEY_ORDER.index(lane_key)
    except ValueError:
        group_index = len(AIR_SPLIT_KEY_ORDER)
    parsed_lane_number = lane_number_value(lane_number) or 0
    return group_index, parsed_lane_number


AIR_SPLIT_LAX_ORIGIN_PORTS = frozenset({"CAN", "CTU", "HKG", "SZX", "XMN"})
AIR_SPLIT_ORD_ORIGIN_PORTS_EXCLUDED = frozenset({"SZX", "XMN"})


def should_split_air_validity(
    origin_country: object,
    destination_country: object,
    origin_port: object,
    destination_port: object,
) -> bool:
    origin_cc = destination_country_code(origin_country)
    dest_cc = destination_country_code(destination_country)
    origin_port_code = normalize_port_name(origin_port)
    dest_port_code = normalize_port_name(destination_port)
    if dest_cc != "US" or origin_cc not in {"CN", "HK"}:
        return False
    if dest_port_code == "LAX":
        return origin_port_code in AIR_SPLIT_LAX_ORIGIN_PORTS
    if dest_port_code == "ORD":
        return origin_port_code not in AIR_SPLIT_ORD_ORIGIN_PORTS_EXCLUDED
    return False


def lane_number_value(value: object) -> int | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    parsed = pd.to_numeric(str(value).replace(" ", ""), errors="coerce")
    if pd.isna(parsed):
        return None
    return int(parsed)


def row_needs_air_validity_extension(
    period_start: date,
    period_end: date,
    update_start: date,
    update_end: date,
) -> bool:
    if period_end is None or update_start is None or update_end is None:
        return False
    return period_end + timedelta(days=1) == update_start


def apply_air_validity_extension(
    values: list[object],
    *,
    columns: dict[str, int],
    valid_from_column: str,
    valid_to_column: str,
    lane_number_column: str | None,
    update_end: date,
    split_validity: bool,
    update_start: date,
    next_lane_number: int | None,
) -> tuple[list[object], list[object], int | None, int]:
    """Return updated row, optional appended split row, next lane number, changed cell count."""
    changed_cells = 0
    row_copy = list(values)

    if split_validity:
        split_row = list(values)
        split_row[columns[valid_from_column] - 1] = format_ra_date(update_start)
        split_row[columns[valid_to_column] - 1] = format_ra_date(update_end)
        return row_copy, [split_row], next_lane_number, changed_cells

    old_valid_to = row_copy[columns[valid_to_column] - 1]
    row_copy[columns[valid_to_column] - 1] = format_ra_date(update_end)
    if str(old_valid_to).strip() != format_ra_date(update_end):
        changed_cells += 1
    return row_copy, [], next_lane_number, changed_cells


def global_update_window(update_lookup: dict[str, pd.Series]) -> tuple[date | None, date | None]:
    starts: list[date] = []
    ends: list[date] = []
    for update_row in update_lookup.values():
        update_start, update_end = update_validity_dates(update_row)
        if update_start is not None and update_end is not None:
            starts.append(update_start)
            ends.append(update_end)
    if not starts:
        return None, None
    return min(starts), max(ends)


def process_air_rate_card_sheet(ws: Worksheet, update_lookup: dict[str, pd.Series]) -> dict[str, int]:
    header_row = _find_rate_card_header_row_from_sheet(ws)
    first_data_row = header_row + 1
    max_col = ws.max_column
    columns = header_column_map(ws, header_row)
    identity_columns = resolve_rate_card_identity_columns(columns)
    lane_columns = resolve_header_lane_columns(columns)
    cost_columns = parse_cost_columns(ws)
    shipment_last_col = shipment_details_last_column(cost_columns)

    valid_from_column = identity_columns["valid_from"]
    valid_to_column = identity_columns["valid_to"]
    lane_number_column = identity_columns.get("lane_number")
    origin_country_column = lane_columns["origin_country"]
    destination_country_column = lane_columns["destination_country"]
    origin_port_column = lane_columns["origin_port"]
    destination_port_column = lane_columns["destination_port"]

    data_rows: list[list[object]] = []
    max_lane_number = 0
    for row_idx in range(first_data_row, ws.max_row + 1):
        values = row_values(ws, row_idx, max_col)
        if all(value is None for value in values[:10]):
            continue
        data_rows.append(values)
        if lane_number_column is not None:
            lane_number = lane_number_value(values[columns[lane_number_column] - 1])
            if lane_number is not None:
                max_lane_number = max(max_lane_number, lane_number)

    new_data_rows: list[list[object]] = []
    pending_split_rows: list[tuple[tuple[int, int], list[object]]] = []
    stats = {
        "lanes_split": 0,
        "rows_created": 0,
        "rows_removed": 0,
        "cost_cells_changed": 0,
        "validity_cells_changed": 0,
        "lanes_added": 0,
        "rows_added_unmapped": 0,
    }
    updated_keys: set[str] = set()
    next_lane_number = max_lane_number + 1 if max_lane_number else None
    global_update_start, global_update_end = global_update_window(update_lookup)

    for values in data_rows:
        lane_key = lane_key_from_values(values, columns, lane_columns)
        update_row = update_lookup.get(lane_key) if lane_key else None
        update_start, update_end = update_validity_dates(update_row) if update_row is not None else (None, None)
        if update_start is None or update_end is None:
            update_start, update_end = global_update_start, global_update_end
        if update_start is None or update_end is None:
            new_data_rows.append(values)
            continue

        period_start = parse_date(values[columns[valid_from_column] - 1])
        period_end = parse_date(values[columns[valid_to_column] - 1])
        if period_start is None or period_end is None:
            new_data_rows.append(values)
            continue

        if not row_needs_air_validity_extension(period_start, period_end, update_start, update_end):
            new_data_rows.append(values)
            continue

        split_validity = (
            lane_key is not None
            and update_row is not None
            and should_split_air_validity(
                values[columns[origin_country_column] - 1],
                values[columns[destination_country_column] - 1],
                values[columns[origin_port_column] - 1],
                values[columns[destination_port_column] - 1],
            )
        )
        updated_row, split_rows, next_lane_number, validity_changes = apply_air_validity_extension(
            values,
            columns=columns,
            valid_from_column=valid_from_column,
            valid_to_column=valid_to_column,
            lane_number_column=lane_number_column,
            update_end=update_end,
            split_validity=split_validity,
            update_start=update_start,
            next_lane_number=next_lane_number,
        )
        if split_rows and update_row is not None and lane_key is not None:
            source_lane_number = (
                values[columns[lane_number_column] - 1] if lane_number_column is not None else None
            )
            for split_row in split_rows:
                pending_split_rows.append(
                    (
                        air_split_sort_key(str(lane_key), source_lane_number),
                        split_row,
                    )
                )
        new_data_rows.append(updated_row)
        if validity_changes:
            stats["validity_cells_changed"] += validity_changes
        if split_rows:
            stats["rows_created"] += len(split_rows)
        if lane_key and update_row is not None and lane_key not in updated_keys:
            updated_keys.add(lane_key)
            stats["lanes_split"] += 1

    pending_split_rows.sort(key=lambda item: item[0])
    appended_rows: list[list[object]] = []
    for _, split_row in pending_split_rows:
        if next_lane_number is not None and lane_number_column is not None:
            split_row[columns[lane_number_column] - 1] = next_lane_number
            next_lane_number += 1
        appended_rows.append(split_row)

    new_data_rows.extend(appended_rows)

    existing_row_count = len(new_data_rows)
    new_data_rows, next_lane_number, add_stats = append_unmapped_lanes(
        new_data_rows,
        update_lookup,
        columns=columns,
        lane_columns=lane_columns,
        identity_columns=identity_columns,
        next_lane_number=next_lane_number,
    )
    stats["lanes_added"] = add_stats["lanes_added"]
    stats["rows_added_unmapped"] = add_stats["rows_added_unmapped"]

    for row_offset, values in enumerate(new_data_rows):
        write_row_values(ws, first_data_row + row_offset, values)
    for row_idx in range(first_data_row + len(new_data_rows), ws.max_row + 1):
        for col in range(1, max_col + 1):
            ws.cell(row_idx, col).value = None

    valid_from_col = columns[valid_from_column]
    valid_to_col = columns[valid_to_column]
    for row_offset, values in enumerate(new_data_rows):
        row_idx = first_data_row + row_offset
        if row_offset >= existing_row_count:
            continue
        lane_key = lane_key_from_values(values, columns, lane_columns)
        if lane_key in update_lookup:
            update_row = update_lookup[lane_key]
            highlight_shipment_details_block(ws, row_idx, shipment_last_col)
            stats["cost_cells_changed"] += apply_cost_updates(
                ws,
                row_idx,
                update_row,
                cost_columns,
            )
            update_start, update_end = update_validity_dates(update_row)
            period_end = parse_date(values[columns[valid_to_column] - 1])
            period_start = parse_date(values[columns[valid_from_column] - 1])
            if (
                update_start is not None
                and update_end is not None
                and period_start is not None
                and period_end is not None
                and (
                    period_end == update_end
                    or (period_start == update_start and period_end == update_end)
                )
            ):
                highlight_cell(ws, row_idx, valid_from_col)
                highlight_cell(ws, row_idx, valid_to_col)
            continue

        period_end = parse_date(values[columns[valid_to_column] - 1])
        period_start = parse_date(values[columns[valid_from_column] - 1])
        if (
            period_start is None
            or period_end is None
            or global_update_end is None
            or period_end != global_update_end
        ):
            continue
        highlight_shipment_details_block(ws, row_idx, shipment_last_col)
        highlight_cell(ws, row_idx, valid_from_col)
        highlight_cell(ws, row_idx, valid_to_col)

    for row_offset in range(existing_row_count, len(new_data_rows)):
        row_idx = first_data_row + row_offset
        values = new_data_rows[row_offset]
        lane_key = lane_key_from_values(values, columns, lane_columns)
        update_row = update_lookup.get(lane_key) if lane_key else None
        if update_row is None:
            continue
        apply_cost_updates(ws, row_idx, update_row, cost_columns, highlight=False)
        highlight_new_lane_row(
            ws,
            row_idx,
            shipment_last_col=shipment_last_col,
            cost_columns=cost_columns,
            valid_from_col=valid_from_col,
            valid_to_col=valid_to_col,
        )

    return stats


def process_sea_cost_only_rate_card_sheet(
    ws: Worksheet,
    update_lookup: dict[str, pd.Series],
) -> dict[str, int]:
    """Update Sea rate cards that have no validity columns (e.g. disabled surcharge layouts)."""
    header_row = _find_rate_card_header_row_from_sheet(ws)
    first_data_row = header_row + 1
    max_col = ws.max_column
    columns = header_column_map(ws, header_row)
    identity_columns = resolve_rate_card_identity_columns(columns, require_validity=False)
    lane_columns = resolve_header_lane_columns(columns)
    cost_columns = parse_cost_columns(ws)
    shipment_last_col = shipment_details_last_column(cost_columns)
    lane_number_column = identity_columns.get("lane_number")

    data_rows: list[list[object]] = []
    max_lane_number = 0
    for row_idx in range(first_data_row, ws.max_row + 1):
        values = row_values(ws, row_idx, max_col)
        if all(value is None for value in values[:10]):
            continue
        data_rows.append(values)
        if lane_number_column is not None:
            lane_number = lane_number_value(values[columns[lane_number_column] - 1])
            if lane_number is not None:
                max_lane_number = max(max_lane_number, lane_number)

    stats = {
        "lanes_split": 0,
        "rows_created": 0,
        "rows_removed": 0,
        "cost_cells_changed": 0,
        "lanes_added": 0,
        "rows_added_unmapped": 0,
    }

    next_lane_number = max_lane_number + 1 if max_lane_number else None
    existing_row_count = len(data_rows)
    new_data_rows, next_lane_number, add_stats = append_unmapped_lanes(
        data_rows,
        update_lookup,
        columns=columns,
        lane_columns=lane_columns,
        identity_columns=identity_columns,
        next_lane_number=next_lane_number,
    )
    stats["lanes_added"] = add_stats["lanes_added"]
    stats["rows_added_unmapped"] = add_stats["rows_added_unmapped"]

    for row_offset, values in enumerate(new_data_rows):
        write_row_values(ws, first_data_row + row_offset, values)
    for row_idx in range(first_data_row + len(new_data_rows), ws.max_row + 1):
        for col in range(1, max_col + 1):
            ws.cell(row_idx, col).value = None

    valid_from_column = identity_columns.get("valid_from")
    valid_to_column = identity_columns.get("valid_to")
    valid_from_col = columns[valid_from_column] if valid_from_column else None
    valid_to_col = columns[valid_to_column] if valid_to_column else None

    for row_offset, values in enumerate(new_data_rows):
        row_idx = first_data_row + row_offset
        if row_offset >= existing_row_count:
            continue
        lane_key = lane_key_from_values(values, columns, lane_columns)
        update_row = update_lookup.get(lane_key) if lane_key else None
        if update_row is None:
            continue
        highlight_shipment_details_block(ws, row_idx, shipment_last_col)
        stats["cost_cells_changed"] += apply_cost_updates(ws, row_idx, update_row, cost_columns)

    for row_offset in range(existing_row_count, len(new_data_rows)):
        row_idx = first_data_row + row_offset
        values = new_data_rows[row_offset]
        lane_key = lane_key_from_values(values, columns, lane_columns)
        update_row = update_lookup.get(lane_key) if lane_key else None
        if update_row is None:
            continue
        stats["cost_cells_changed"] += apply_cost_updates(
            ws, row_idx, update_row, cost_columns, highlight=False
        )
        highlight_new_lane_row(
            ws,
            row_idx,
            shipment_last_col=shipment_last_col,
            cost_columns=cost_columns,
            valid_from_col=valid_from_col,
            valid_to_col=valid_to_col,
        )

    return stats


def process_rate_card_sheet(ws: Worksheet, update_lookup: dict[str, pd.Series]) -> dict[str, int]:
    if is_air_rate_card(ws):
        return process_air_rate_card_sheet(ws, update_lookup)

    header_row = _find_rate_card_header_row_from_sheet(ws)
    first_data_row = header_row + 1
    max_col = ws.max_column
    columns = header_column_map(ws, header_row)
    identity_columns = resolve_rate_card_identity_columns(columns, require_validity=False)
    if "valid_from" not in identity_columns or "valid_to" not in identity_columns:
        return process_sea_cost_only_rate_card_sheet(ws, update_lookup)

    lane_columns = resolve_header_lane_columns(columns)
    cost_columns = parse_cost_columns(ws)
    shipment_last_col = shipment_details_last_column(cost_columns)

    lane_id_column = identity_columns["lane_id"]
    valid_from_column = identity_columns["valid_from"]
    valid_to_column = identity_columns["valid_to"]
    lane_number_column = identity_columns.get("lane_number")

    data_rows: list[list[object]] = []
    lane_groups: dict[int, list[tuple[int, list[object], tuple[date, date]]]] = {}

    for row_idx in range(first_data_row, ws.max_row + 1):
        values = row_values(ws, row_idx, max_col)
        if all(value is None for value in values[:10]):
            continue
        data_rows.append(values)
        lane_id = lane_id_value(values[columns[lane_id_column] - 1])
        period_start = parse_date(values[columns[valid_from_column] - 1])
        period_end = parse_date(values[columns[valid_to_column] - 1])
        if lane_id is None or period_start is None or period_end is None:
            continue
        lane_groups.setdefault(lane_id, []).append((len(data_rows) - 1, values, (period_start, period_end)))

    processed_lane_ids: set[int] = set()
    new_data_rows: list[list[object]] = []
    stats = {
        "lanes_split": 0,
        "rows_created": 0,
        "rows_removed": 0,
        "cost_cells_changed": 0,
        "lanes_added": 0,
        "rows_added_unmapped": 0,
    }
    max_lane_number = 0

    for values in data_rows:
        if lane_number_column is not None:
            lane_number = lane_number_value(values[columns[lane_number_column] - 1])
            if lane_number is not None:
                max_lane_number = max(max_lane_number, lane_number)
        lane_id = lane_id_value(values[columns[lane_id_column] - 1])

        if lane_id is None or lane_id not in lane_groups:
            new_data_rows.append(values)
            continue

        if lane_id in processed_lane_ids:
            continue

        processed_lane_ids.add(lane_id)
        group = lane_groups[lane_id]
        template_values = group[0][1]
        lane_key = lane_key_from_values(template_values, columns, lane_columns)
        update_row = update_lookup.get(lane_key)
        if update_row is None:
            for _, row_values_list, _ in group:
                new_data_rows.append(row_values_list)
            continue

        update_start, update_end = update_validity_dates(update_row)
        if update_start is None or update_end is None:
            for _, row_values_list, _ in group:
                new_data_rows.append(row_values_list)
            continue

        ra_periods = [period for _, _, period in group]
        lane_templates = [(period, values) for _, values, period in group]
        segments = split_validity_periods(update_start, update_end, ra_periods)
        stats["lanes_split"] += 1
        stats["rows_removed"] += len(group)
        base_lane_number = (
            group[0][1][columns[lane_number_column] - 1]
            if lane_number_column is not None
            else None
        )

        for segment in segments:
            row_copy = choose_template_row(segment, lane_templates)
            row_copy[columns[valid_from_column] - 1] = format_ra_date(segment[0])
            row_copy[columns[valid_to_column] - 1] = format_ra_date(segment[1])
            if lane_number_column is not None and base_lane_number is not None:
                row_copy[columns[lane_number_column] - 1] = base_lane_number
            new_data_rows.append(row_copy)
            stats["rows_created"] += 1

    next_lane_number = max_lane_number + 1 if max_lane_number else None
    existing_row_count = len(new_data_rows)
    new_data_rows, next_lane_number, add_stats = append_unmapped_lanes(
        new_data_rows,
        update_lookup,
        columns=columns,
        lane_columns=lane_columns,
        identity_columns=identity_columns,
        next_lane_number=next_lane_number,
    )
    stats["lanes_added"] = add_stats["lanes_added"]
    stats["rows_added_unmapped"] = add_stats["rows_added_unmapped"]

    for row_offset, values in enumerate(new_data_rows):
        write_row_values(ws, first_data_row + row_offset, values)
    for row_idx in range(first_data_row + len(new_data_rows), ws.max_row + 1):
        for col in range(1, max_col + 1):
            ws.cell(row_idx, col).value = None

    for row_offset, values in enumerate(new_data_rows):
        row_idx = first_data_row + row_offset
        if row_offset >= existing_row_count:
            continue
        lane_id = lane_id_value(values[columns[lane_id_column] - 1])
        if lane_id is None or lane_id not in processed_lane_ids:
            continue
        template_values = lane_groups[lane_id][0][1]
        lane_key = lane_key_from_values(template_values, columns, lane_columns)
        update_row = update_lookup.get(lane_key)
        if update_row is None:
            continue
        highlight_shipment_details_block(ws, row_idx, shipment_last_col)
        stats["cost_cells_changed"] += apply_cost_updates(ws, row_idx, update_row, cost_columns)

    valid_from_col = columns[valid_from_column]
    valid_to_col = columns[valid_to_column]
    for row_offset in range(existing_row_count, len(new_data_rows)):
        row_idx = first_data_row + row_offset
        values = new_data_rows[row_offset]
        lane_key = lane_key_from_values(values, columns, lane_columns)
        update_row = update_lookup.get(lane_key) if lane_key else None
        if update_row is None:
            continue
        stats["cost_cells_changed"] += apply_cost_updates(
            ws, row_idx, update_row, cost_columns, highlight=False
        )
        highlight_new_lane_row(
            ws,
            row_idx,
            shipment_last_col=shipment_last_col,
            cost_columns=cost_columns,
            valid_from_col=valid_from_col,
            valid_to_col=valid_to_col,
        )

    return stats


def _find_rate_card_header_row_from_sheet(ws: Worksheet) -> int:
    for row_idx in range(1, min(ws.max_row, 40) + 1):
        values = {
            str(ws.cell(row_idx, col).value).strip().upper()
            for col in range(1, ws.max_column + 1)
            if ws.cell(row_idx, col).value is not None and str(ws.cell(row_idx, col).value).strip()
        }
        if "LANE #" in values and DESTINATION_COUNTRY_COLUMN in values:
            return row_idx
    raise ValueError("Could not find Rate card header row.")


def apply_update_to_workbook(source_path: Path, output_path: Path, update_lookup: dict[str, pd.Series]) -> dict[str, int]:
    workbook = load_workbook(source_path)
    if RATE_CARD_SHEET not in workbook.sheetnames:
        raise ValueError(f"Sheet '{RATE_CARD_SHEET}' not found in {source_path.name}")
    stats = process_rate_card_sheet(workbook[RATE_CARD_SHEET], update_lookup)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return stats


def run_apply_rate_card_update(
    *,
    previous_ra_paths: list[Path] | None = None,
    update_path: Path | None = None,
) -> list[Path]:
    print("\n--- Apply update to Rate card layout ---")
    update = load_update_rows(update_path)
    update_lookup = build_update_lookup(update)
    print(f"  Update lanes loaded: {len(update_lookup):,}")

    if previous_ra_paths is None:
        previous_ra_paths = sorted(
            path for path in PREVIOUS_RA_DIR.glob("*.xlsx") if not path.name.startswith("~$")
        )

    saved_paths: list[Path] = []
    for source_path in previous_ra_paths:
        output_path = OUTPUT_DIR / f"Updated - {source_path.name}"
        stats = apply_update_to_workbook(source_path, output_path, update_lookup)
        print(
            f"  {source_path.name}: split {stats['lanes_split']} lane(s), "
            f"created {stats['rows_created']} row(s), "
            f"added {stats.get('lanes_added', 0)} new lane(s), "
            f"changed {stats['cost_cells_changed']} cost cell(s) -> {output_path.name}"
        )
        saved_paths.append(output_path)
    return saved_paths


def main() -> None:
    from paths import configure_colab_drive, ensure_data_dirs, print_paths

    configure_colab_drive()
    ensure_data_dirs()
    print_paths()
    run_apply_rate_card_update()


if __name__ == "__main__":
    main()
